from openpyxl import load_workbook
import pandas as pd

from src.pipeline import run_pipeline
from src.config import AppConfig
from src.reporting import (
    generate_excel_report,
    create_summary,
)


def test_generate_excel_report(sample_df, sample_config, tmp_path):

    results = run_pipeline(sample_df, sample_config)

    output_path = tmp_path / "new" / "nested" / "data_quality_report.xlsx"

    generate_excel_report(results, output_path)

    assert output_path.exists()


    workbook = load_workbook(output_path)

    assert "Summary" in workbook.sheetnames
    assert "Missing Values" in workbook.sheetnames
    assert "Duplicates" in workbook.sheetnames
    assert "Invalid Values" in workbook.sheetnames
    assert "Invalid Dates" in workbook.sheetnames
    assert "Data Types" in workbook.sheetnames
    assert "Numeric Profiling" in workbook.sheetnames
    assert "Categorical Profiling" in workbook.sheetnames
    assert list(workbook["Dataset Overview"].values)[1][:3] == ("rows", 1010, 1000)
    missing = list(workbook["Missing Values"].values)
    revenue_before = next(row for row in missing[1:] if row[:2] == ("BEFORE", "revenue"))
    assert revenue_before[2:4] == (20, 1010)
    assert abs(revenue_before[4] - 20 / 1010 * 100) < 1e-12
    date_after = next(row for row in missing[1:] if row[:2] == ("AFTER", "date"))
    assert date_after[2:4] == (1, 1000)
    for sheet in workbook:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == sheet.dimensions
        assert sheet["A1"].font.bold
        assert sheet.column_dimensions["A"].width >= 14
        assert not any(cell.data_type == "e" for row in sheet for cell in row)
    assert workbook["Missing Values"]["E2"].number_format == '0.00"%"'
    assert workbook["Numeric Profiling"]["D2"].number_format == "#,##0.00"
    workbook.close()


def test_create_summary(sample_df, sample_config):

    results = run_pipeline(sample_df, sample_config)

    summary = create_summary(results)

    assert summary.loc["Missing Values", "Before"] == 20
    assert summary.loc["Missing Values", "After"] == 1
    assert summary.loc["Duplicate Rows", "Before"] == 10
    assert summary.loc["Duplicate Rows", "After"] == 0
    assert summary.loc["Invalid Dates", "Before"] == 1
    assert summary.loc["Invalid Values", "After"] == 3


def test_source_text_is_not_excel_formula(tmp_path):
    results = run_pipeline(pd.DataFrame({"=column": ["=1+1"]}), AppConfig())
    output = tmp_path / "literal.xlsx"
    generate_excel_report(results, output)
    workbook = load_workbook(output)
    categorical = workbook["Categorical Profiling"]
    assert categorical["B2"].value == "=column"
    assert categorical["F2"].value == "=1+1"
    assert categorical["B2"].data_type == categorical["F2"].data_type == "s"
    workbook.close()
