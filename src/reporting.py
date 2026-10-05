"""Excel output: explicit stages, defined metrics, and lightweight formatting."""
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def create_summary(results):
    """Counts are not additive: rule violations may overlap within a row."""
    rows = {}
    for title, key, field, unit in [
        ("Missing Values", "missing_values", "missing_count", "missing cells"),
        ("Duplicate Rows", "duplicates", "duplicate_count", "extra exact duplicate rows"),
        ("Invalid Values", "invalid_values", "invalid_count", "rule violations, not distinct rows"),
        ("Invalid Dates", "invalid_dates", "invalid_count", "nonmissing unparseable date cells"),
    ]:
        before, after = results[key], results[f"cleaned_{key}"]
        rows[title] = {
            "Before": int(before[field] if isinstance(before, dict) else before[field].sum()),
            "After": int(after[field] if isinstance(after, dict) else after[field].sum()),
            "Unit": unit,
        }
        status = "checked"
        if isinstance(before, pd.DataFrame) and "status" in before:
            if before.empty:
                status = "not configured"
            elif before["status"].eq("missing_column").all():
                status = "unchecked: missing configured columns"
            elif before["status"].eq("missing_column").any():
                status = "partial: missing configured columns"
        if status == "checked" and results["dataset_profile"]["rows"] == 0:
            status = "no rows"
        rows[title]["Assessment"] = status
    return pd.DataFrame.from_dict(rows, orient="index").rename_axis("metric")


def _stages(results, key):
    """Stack detail tables rather than hiding the after-cleaning measurements."""
    frames = []
    for stage, prefix in [("BEFORE", ""), ("AFTER", "cleaned_")]:
        value = results[prefix + key]
        frame = pd.DataFrame([value]) if isinstance(value, dict) else value.copy()
        if frame.index.name == "column":
            frame = frame.reset_index()
        frame.insert(0, "stage", stage)
        frames.append(frame)
    return pd.concat(frames, ignore_index=True)


def generate_excel_report(results, output_path, context=None):
    """Create parent directories and write tables with readable headers and units."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    overview = pd.DataFrame({
        "Before": results["dataset_profile"],
        "After": results["cleaned_dataset_profile"],
    }).rename_axis("metric").reset_index()
    notes = [
        ("Stages", "BEFORE is the input; AFTER follows duplicate removal, date parsing, text trimming, and imputation."),
        ("Missing values", "Missing cells per column; percentage denominator = rows in that stage."),
        ("Duplicates", "Extra exact duplicate rows, keeping first; percentage denominator = rows in that stage."),
        ("Numeric rules", "Finite numeric strings are evaluated numerically without changing source data. Non-numeric/infinite values count once as numeric_type."),
        ("Rule denominators", "evaluated_count = nonmissing values; range rules exclude numeric_type failures. Missing columns are unchecked, not passed."),
        ("Invalid values", "Sum of per-rule violations; a row can fail several rules. Nulls are counted only as missing."),
        ("Invalid dates", "Nonmissing values that fail the declared format; denominator = nonmissing date values. Coercion creates missing values, not corrected dates."),
        ("Date convention", "Declared format; timestamps normalized to UTC. Naive dates are treated as UTC. No mixed-format inference or numeric epoch parsing."),
        ("Imputation", "Finite numeric / nonmissing categorical donors only; identifiers can use skip. Datetimes/booleans and all-null columns are not imputed. Mode ties use pandas' first mode."),
        ("Analytical caution", "Imputation uses all observed values, including flagged business values, and can change distributions. Review validity before using outputs for analysis."),
        ("Zero denominators", "Percentages are 0 when nothing is evaluated; consult row/evaluated counts and status."),
        ("Data types", "Observed pandas dtypes, not validation against an expected schema."),
        ("Profiling", "Numeric statistics describe all observed numeric values, including identifiers/flagged values. most_common is the first mode in ties."),
        ("Residual duplicates", "Only exact input duplicates are removed. Normalization/imputation can create duplicates; AFTER checks keep them visible."),
        ("No quality score", "Metrics have different units and may overlap; they are not combined into a score."),
    ]
    tables = {
        "Summary": create_summary(results).reset_index(),
        "Dataset Overview": overview,
        "Cleaning Actions": results["actions"],
        "Configured Columns": results["configured_columns"],
        "Missing Values": _stages(results, "missing_values"),
        "Duplicates": _stages(results, "duplicates"),
        "Invalid Values": _stages(results, "invalid_values"),
        "Invalid Dates": _stages(results, "invalid_dates"),
        "Data Types": _stages(results, "data_types"),
        "Numeric Profiling": _stages(results, "numeric_profile"),
        "Categorical Profiling": _stages(results, "categorical_profile"),
        "Notes": pd.DataFrame(notes, columns=["topic", "definition"]),
        "Run Configuration": pd.DataFrame(
            [(key, json.dumps(value, ensure_ascii=False)) for key, value in results["config"].items()],
            columns=["setting", "value"],
        ),
    }
    if context:
        tables["Run Context"] = pd.DataFrame(context.items(), columns=["item", "value"])
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, frame in tables.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 30
            for index, header in enumerate(frame.columns, start=1):
                cells = list(sheet.iter_cols(min_col=index, max_col=index))[0]
                width = min(65, max(14, max(len(str(cell.value or "")) for cell in cells) + 2))
                sheet.column_dimensions[get_column_letter(index)].width = width
                for cell in cells[1:]:
                    cell.font = Font(name="Arial", size=10)
                    # Source labels/categories are literal text, even when they start with '='.
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if isinstance(cell.value, (int, float)):
                        # Values already use a 0-100 scale; '%' must be literal.
                        if "percentage" in str(header):
                            cell.number_format = '0.00"%"'
                        elif header in {"mean", "std", "min", "25%", "50%", "75%", "max"}:
                            cell.number_format = "#,##0.00"
                        else:
                            cell.number_format = "#,##0"
            for row in sheet.iter_rows(min_row=2):
                longest = max((len(str(cell.value or "")) / sheet.column_dimensions[cell.column_letter].width
                               for cell in row), default=1)
                sheet.row_dimensions[row[0].row].height = min(100, max(18, (int(longest) + 1) * 16))
